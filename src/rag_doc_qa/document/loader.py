"""文档加载：PDF / Word / Markdown。"""
from __future__ import annotations

from pathlib import Path
from typing import List

from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


def load_pdf(path: Path) -> List[Document]:
    """使用 PyMuPDF 解析 PDF（对中文支持较好）。"""
    import fitz  # PyMuPDF

    docs: List[Document] = []
    with fitz.open(path) as pdf:
        for i, page in enumerate(pdf):
            text = page.get_text("text")
            if not text.strip():
                continue
            docs.append(
                Document(page_content=text, metadata={"source": str(path), "page": i + 1})
            )
    return docs


def load_docx(path: Path) -> List[Document]:
    """使用 python-docx 解析 Word 文档（含表格文本）。"""
    import docx

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return [
        Document(page_content="\n".join(paragraphs), metadata={"source": str(path), "page": 1})
    ]


def load_markdown(path: Path) -> List[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [Document(page_content=text, metadata={"source": str(path), "page": 1})]


def load_document(path: str | Path) -> List[Document]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in {".md", ".markdown", ".txt"}:
        return load_markdown(path)
    raise ValueError(f"不支持的文件类型: {ext}（支持 {sorted(SUPPORTED_EXTENSIONS)}）")


def load_documents(paths) -> List[Document]:
    """加载一批文件或目录，返回 Document 列表。"""
    docs: List[Document] = []
    for p in paths:
        path = Path(p)
        if path.is_dir():
            for f in sorted(path.rglob("*")):
                if f.suffix.lower() in SUPPORTED_EXTENSIONS:
                    docs.extend(load_document(f))
        else:
            docs.extend(load_document(path))
    return docs
